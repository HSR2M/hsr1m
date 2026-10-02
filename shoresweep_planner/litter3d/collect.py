"""지도 위 쓰레기(파이프라인 탐지 위치 + 3D 부피, 또는 업체 라벨) + 추정 무게 + 지형 → 작업자용 최적 수거 계획.

입력: LitterObject 목록 (pipeline_io 또는 load_geojson 이 만든다)
      (선택) terrain.Terrain — 정사영상 또는 DEM 에서 만든 물·숲·맨땅 비용 격자
출력: 구역 묶기 → 출발지 기준 최적 순회 → 구역별 마대·시간·주의 → 운반 방식(현장 적치/들고 이동) → 일차 분할

무게 (basis 별)
  "3d"    3D 부피(높이지도) × 겉보기 밀도 (classes.py). 부피 측정 신뢰도에 따라 부피 범위를 넓힌다
  "proxy" 부피 측정이 실패한 물체 → 같은 종류의 측정 부피 중앙값(신뢰 높음·중간 우선) 으로 대체, 범위 ×0.5–2
  "2d"    면적만 있을 때 (업체 라벨·정사영상 탐지): 면적 × 채움률 × 두께 × 겉보기 밀도 (채움률·두께는 가정값)
  "agg"   파이프라인이 이미 무게를 낸 정거장(stop) 묶음 (물체가 수백 개 넘는 큰 현장): 종류별 kg 범위를 그대로 쓰고 보정 계수만 곱함
  "count" 크기 정보가 전혀 없을 때: 개수 × 종류별 평균무게 (Andriolo 2024)

모델 요약
  거리   지형이 있으면 8방향 최단경로의 '걷기 환산 m' (물 통행 불가·숲 ×3·보트 모드는 물 ×0.5), 없으면 직선 × 우회 배수
  구역   걷기 환산 거리 link_m 안의 물체를 한 구역으로 (single-linkage)
  순서   objective="distance": 구역 순회 최단 (최근접 이웃 + 2-opt + or-opt, 8개 이하 완전탐색)
         objective="weight"  : 다음 구역 = kg ÷ (이동분 + 작업분) 이 가장 큰 곳 (무거운 곳부터)
  운반   carry="pile" : 마대를 구역에 모아 두고 나중에 회수 (적재 제한 없음)
         carry="carry": 들고 이동. 팀 적재량(인원 × 1인 kg, 인원 × 1인 마대 수) 넘으면 출발지로 돌아왔다가 계속
  시간   이동 = 환산 m ÷ 걷기 속도 × 짐 감속, 작업 = (물체당 기본분 × 개수 + m² 당 분) ÷ 인원 + 무거운 물체 추가분
  일정   하루 작업시간을 넘으면 다음 일차

⚠ 출처 있는 값: 1인 들기 한계 23 kg (NIOSH), 스티로폼 밀도 11–32 kg/m³ (classes.py)
⚠ 가정값: 겉보기 밀도 대부분, 부피 측정 신뢰도별 범위, 2D 라벨 채움률·두께, 작업 속도, 걷기 속도, 우회 배수, 지형 통행 배수,
   마대·운반 적재량. 모두 CollectParams / MATERIALS / VOL_RANGE / terrain.COST 에 모여 있고 결과에 '가정값' 으로 표기된다.
"""
from __future__ import annotations

import json
import math
from dataclasses import dataclass, field, asdict
from pathlib import Path

import numpy as np

from .classes import CLASSES, KO_NAMES, map_class
from .plan import NIOSH_LIFT_LIMIT_KG, PlanParams

# ─────────────────────────── 재질(클래스) → 2D 추정 규칙·표시 (가정값) ───────────────────────────
@dataclass(frozen=True)
class Material:
    code: str            # 클래스 키 (classes.CLASSES)
    class_name: str
    ko: str
    color: str
    fill: float          # 사각형 라벨 안에서 실제 물체가 차지하는 비율 (2D 추정용)
    t_min: float         # 두께(m) (2D 추정용)
    t_typ: float
    t_max: float
    handling_note: str = ""
    tool: str = ""
    default_vol_m3: float = 0.01   # 측정값이 하나도 없을 때 대체 부피 (가정값)


def _m(code, color, fill, t, note="", tool="", dv=0.01):
    return Material(code, code, KO_NAMES.get(code, code), color, fill, t[0], t[1], t[2], note, tool, dv)


MATERIALS: dict[str, Material] = {
    "styrofoam_fragment": _m("styrofoam_fragment", "#ffd60a", 0.5, (0.03, 0.10, 0.25),
                             "가볍지만 부피가 큼 → 마대가 먼저 참. 큰 덩어리는 부수지 말고 통째로 담기(미세조각 방지)", dv=0.01),
    "styrofoam_buoy": _m("styrofoam_buoy", "#ffc300", 0.6, (0.15, 0.30, 0.40),
                         "부피가 큼 → 마대보다 톤백·그물망. 부서지면 미세조각이 흩어지므로 통째로 운반", dv=0.03),
    "styrofoam_box": _m("styrofoam_box", "#ffe066", 0.6, (0.20, 0.30, 0.35), "비어 있으면 가볍지만 부피가 큼. 겹쳐 담기", dv=0.04),
    "rope": _m("rope", "#ff7a1a", 0.4, (0.03, 0.08, 0.20), "젖으면 매우 무거움. 모래에 묻힌 부분이 있을 수 있음", "칼·가위", dv=0.02),
    "net": _m("net", "#d7263d", 0.4, (0.05, 0.12, 0.30), "모래·돌이 엉켜 있으면 2인 이상. 묻힌 부분은 잘라 내기", "칼·가위·삽", dv=0.03),
    "other_plastic": _m("other_plastic", "#2a78d6", 0.5, (0.02, 0.05, 0.10), "용기·봉지·시트. 날카로운 조각 주의", dv=0.004),
    "pet_bottle": _m("pet_bottle", "#1baf7a", 0.6, (0.06, 0.08, 0.10), "밟아 찌그러뜨려 부피 줄이기", dv=0.0012),
    "plastic_buoy": _m("plastic_buoy", "#0097a7", 0.6, (0.15, 0.30, 0.40), "속이 비어 가볍지만 부피가 큼. 그물망에 모아 운반", dv=0.03),
    "chinese_buoy": _m("chinese_buoy", "#0097a7", 0.6, (0.15, 0.30, 0.40), "속이 비어 가볍지만 부피가 큼", dv=0.03),
    "cardboard": _m("cardboard", "#b5651d", 0.9, (0.10, 0.30, 0.40), "접어서 부피 줄이기. 젖으면 무겁고 찢어짐", dv=0.05),
    "glass": _m("glass", "#8e44ad", 0.7, (0.05, 0.07, 0.10), "깨진 유리 → 두꺼운 장갑, 따로 담기", "두꺼운 장갑", dv=0.001),
    "metal": _m("metal", "#7f8c8d", 0.6, (0.05, 0.10, 0.15), "날카로운 모서리 주의, 따로 담기", "두꺼운 장갑", dv=0.001),
    "wood": _m("wood", "#6d4c41", 0.7, (0.03, 0.05, 0.10), "젖은 유목은 무거움 → 2인 또는 잘라서", "톱", dv=0.02),
    "cloth": _m("cloth", "#e84393", 0.5, (0.03, 0.05, 0.10), "젖으면 무거움", dv=0.005),
    "rubber": _m("rubber", "#2c3e50", 0.7, (0.10, 0.20, 0.25), "타이어는 2인 또는 운반구", "운반구", dv=0.05),
    "ceramic": _m("ceramic", "#95a5a6", 0.8, (0.03, 0.05, 0.10), "무거움 → 작은 양씩", dv=0.002),
    "vegetation": _m("vegetation", "#9acd32", 0.5, (0.02, 0.05, 0.10), "쓰레기 아님"),
    "unknown": _m("unknown", "#a0a8b3", 0.4, (0.02, 0.05, 0.10), "종류 미확인 → 현장에서 확인", dv=0.005),
}
# 업체 재질 코드 (문갑도 GeoJSON) 와 '업체 방식' 무게 계수 (면적 × kg/m²; 실측이 아니라 계산값)
COMPANY_CODES = {"STY": "styrofoam_fragment", "ROP": "rope", "FIS": "net", "PLA": "other_plastic"}
COMPANY_KG_PER_M2 = {"STY": 0.012, "ROP": 0.024, "FIS": 0.024, "PLA": 0.020}
COMPANY_CODE_NAME = {"STY": "스티로폼", "ROP": "로프", "FIS": "어망", "PLA": "플라스틱"}
# 3D 부피 측정 신뢰도 → 부피 범위 배수 (가정값; 선회 촬영 상자 3개 오차 −10~+20 %, 단방향 통과 촬영은 더 큼)
VOL_RANGE = {"높음": (0.8, 1.25), "중간": (0.6, 1.5), "낮음": (0.5, 2.0), "": (0.7, 1.4)}
PROXY_RANGE = (0.5, 2.0)      # 대체 부피(같은 종류 중앙값) 의 범위 (가정값)
COUNT_RANGE = (0.5, 2.0)      # 개수 × 평균무게 의 범위 (가정값)
BASIS_KO = {"3d": "3D 부피 측정", "proxy": "같은 종류 측정값 대체", "2d": "면적 기반(2D)", "agg": "파이프라인 추정(정거장)", "count": "개수 × 평균무게"}


def company_code_of(code: str) -> str:
    """우리 클래스 → 업체 재질 코드 ('업체 방식' 비교용). 업체 분류에 없는 종류는 플라스틱 계수."""
    cls = map_class(code)
    if cls.startswith("styrofoam"):
        return "STY"
    return {"rope": "ROP", "net": "FIS"}.get(cls, "PLA")


def material(code: str) -> Material:
    """클래스 키·업체 코드·탐지 클래스명 무엇을 주든 Material (code 는 우리 클래스 키)."""
    if code in MATERIALS:
        return MATERIALS[code]
    if code in COMPANY_CODES:
        return MATERIALS[COMPANY_CODES[code]]
    return MATERIALS.get(map_class(code), MATERIALS["unknown"])


def materials_js(codes=None) -> dict:
    out = {}
    for c in (codes or MATERIALS.keys()):
        m = material(c); spec = CLASSES.get(m.class_name, CLASSES["unknown"])
        out[m.code] = {"ko": m.ko, "color": m.color, "fill": m.fill, "t": [m.t_min, m.t_typ, m.t_max],
                       "handling": m.handling_note, "tool": m.tool, "rho": [spec.rho_min, spec.rho_typ, spec.rho_max],
                       "wet": spec.wet_factor_max, "company": company_code_of(m.code)}
    return out


# ─────────────────────────── 파라미터 ───────────────────────────
@dataclass
class CollectParams:
    link_m: float = 150.0             # 가정값: 이 걷기거리 안의 물체는 같은 구역
    detour: float = 1.4               # 가정값: 지형 없을 때 직선 → 실제 거리 배수
    walk_kmh: float = 3.0             # 가정값: 짐 들고 걷는 속도
    load_slow: float = 0.004          # 가정값: 1인당 짐 1 kg 마다 걷기 속도 0.4 % 감소 (최대 40 %)
    item_min: float = 2.0             # 가정값: 물체 1개 줍기·담기 기본 시간(분)
    min_per_m2: float = 1.5           # 가정값: 면적 1 m² 당 추가 시간(분)
    heavy_extra_min: float = 5.0      # 가정값: 2인 운반 물체 추가 시간(분)
    workers: int = 2
    hours_per_day: float = 4.0
    round_trip: bool = True
    weight_source: str = "ours"       # "ours" | "company"
    weight_stat: str = "typ"          # ours 일 때 "min" | "typ" | "max"
    travel: str = "walk"              # "walk" | "boat"  (지형 있을 때)
    carry: str = "pile"               # "pile" 현장 적치 | "carry" 들고 이동
    carry_kg_per_person: float = 15.0 # 가정값: 1인이 들고 걷는 무게
    carry_bags_per_person: int = 2    # 가정값: 1인이 들고 걷는 마대 수
    objective: str = "distance"       # "distance" | "weight"
    teams: int = 1                    # 동시에 투입하는 팀 수 (순회를 시간 균형으로 나눔)
    calib: dict | None = None         # 실측 보정 계수 {클래스: 배수}. 우리 추정 무게에 곱함 (예: 현장 마대 저울값으로 산출)
    include_codes: list[str] | None = None   # None 이면 전부
    min_kg: float = 0.0               # 이 무게(계획값) 미만 물체는 건너뜀
    depot_lonlat: tuple[float, float] | None = None
    depot_name: str = "출발지"
    zone_word: str = "해안"           # 구역 이름 ("북쪽 해안 ①" / "북쪽 구역 ①")
    bag: PlanParams = field(default_factory=PlanParams)

    def to_dict(self) -> dict:
        d = asdict(self)
        d["bag"] = self.bag.to_dict()
        d["calib"] = dict(self.calib or {})
        d["_note"] = "link_m·detour·walk_kmh·load_slow·item_min·min_per_m2·heavy_extra_min·운반·마대 적재량은 가정값. 23 kg 는 NIOSH."
        return d


# ─────────────────────────── 데이터 ───────────────────────────
@dataclass
class LitterObject:
    obj_id: str
    seq: int
    code: str                 # 우리 클래스 키
    class_name: str
    class_ko: str
    lon: float
    lat: float
    x_m: float
    y_m: float
    area_m2: float            # 바닥 면적 (3D: 발자국, 2D: 라벨 면적, agg: 합)
    w_m: float                # 가로 (3D: 길이 L)
    h_m: float                # 세로 (3D: 너비 W)
    company_kg: float | None  # 업체 제공값 또는 '업체 방식(면적×고정계수)' 계산값
    kg_min: float
    kg_typ: float
    kg_max: float
    volume_m3: float          # 대표 겉부피 (마대 계산용)
    image_path: str | None
    survey_date: str | None
    basis: str = "2d"         # "3d" | "proxy" | "2d" | "agg" | "count"
    height_m: float = 0.0     # 높이 (3D)
    vol_min_m3: float = 0.0
    vol_max_m3: float = 0.0
    confidence: str = ""      # 3D 측정 신뢰도 문구
    truth_m3: float | None = None   # 정답 부피 (크기를 아는 시험 물체)
    n_items: int = 1          # agg: 묶인 물체 수
    item_kg_max: float = 0.0  # 가장 무거운 낱개 물체의 최대 무게 (2인 운반 판단; 단일 물체면 kg_max)
    kg_by: dict | None = None # agg 혼합: {클래스: [kmin, ktyp, kmax, vol_m3, area_m2, n]}
    members: list | None = None   # agg 상세: [{cls, ko, kg, area}]
    source: str = ""          # 자료 출처 설명 (예: "objects3d.csv #19 3D 높이지도")
    heavy: bool = False
    note: str = ""
    zone: int = -1
    order: int = 0
    included: bool = True

    @property
    def plan_kg(self) -> float:
        return getattr(self, "_plan_kg", self.kg_typ)

    @property
    def dims_cm(self) -> str:
        if self.basis in ("3d", "proxy") and self.height_m:
            return f"{self.w_m * 100:.0f}×{self.h_m * 100:.0f}×{self.height_m * 100:.0f} cm"
        return f"{self.w_m:.1f}×{self.h_m:.1f} m"

    def to_dict(self) -> dict:
        d = asdict(self)
        d["plan_kg"] = round(self.plan_kg, 4)
        d["dims"] = self.dims_cm
        d["basis_ko"] = BASIS_KO.get(self.basis, self.basis)
        return d


@dataclass
class WeightEstimate:
    kg_min: float
    kg_typ: float
    kg_max: float
    vol_min: float
    vol_typ: float
    vol_max: float
    basis: str
    formula: str


def _rho(code: str):
    m = material(code)
    return m, CLASSES.get(m.class_name, CLASSES["unknown"])


def estimate_3d(code: str, vol_m3: float, confidence: str = "") -> WeightEstimate:
    """3D 부피 × 겉보기 밀도. 신뢰도('높음'·'중간'·'낮음…') 에 따라 부피 범위를 넓힌다 (가정값)."""
    m, spec = _rho(code)
    level = next((k for k in ("높음", "중간", "낮음") if confidence.startswith(k)), "")
    lo, hi = VOL_RANGE[level]
    v_min, v_max = vol_m3 * lo, vol_m3 * hi
    kg = (v_min * spec.rho_min, vol_m3 * spec.rho_typ, v_max * spec.rho_max * spec.wet_factor_max)
    formula = (f"3D 부피 {vol_m3 * 1000:.1f} L (신뢰 {level or '보통'}: ×{lo}–{hi}) × ρ {spec.rho_min}/{spec.rho_typ}/{spec.rho_max} kg/m³ "
               f"(최대는 젖음 ×{spec.wet_factor_max})")
    return WeightEstimate(kg[0], kg[1], kg[2], v_min, vol_m3, v_max, "3d", formula)


def estimate_proxy(code: str, vol_m3: float, why: str = "같은 종류 측정 부피 중앙값") -> WeightEstimate:
    m, spec = _rho(code)
    lo, hi = PROXY_RANGE
    kg = (vol_m3 * lo * spec.rho_min, vol_m3 * spec.rho_typ, vol_m3 * hi * spec.rho_max * spec.wet_factor_max)
    formula = f"부피 측정 실패 → {why} {vol_m3 * 1000:.1f} L (×{lo}–{hi}) × ρ {spec.rho_min}/{spec.rho_typ}/{spec.rho_max} kg/m³"
    return WeightEstimate(kg[0], kg[1], kg[2], vol_m3 * lo, vol_m3, vol_m3 * hi, "proxy", formula)


def estimate_2d(code: str, area_m2: float) -> tuple[float, float, float, float, str]:
    """반환 (kg_min, kg_typ, kg_max, volume_typ_m3, 계산식). 업체 사각형 라벨·항공 정사영상 등 면적만 있을 때."""
    m, spec = _rho(code)
    a = area_m2 * m.fill
    v_min, v_typ, v_max = a * m.t_min, a * m.t_typ, a * m.t_max
    kg = (v_min * spec.rho_min, v_typ * spec.rho_typ, v_max * spec.rho_max * spec.wet_factor_max)
    formula = (f"{area_m2:.2f} m² × 채움 {m.fill} × 두께 {m.t_min}/{m.t_typ}/{m.t_max} m × "
               f"ρ {spec.rho_min}/{spec.rho_typ}/{spec.rho_max} kg/m³ (최대는 젖음 ×{spec.wet_factor_max})")
    return kg[0], kg[1], kg[2], v_typ, formula


def estimate_2d_full(code: str, area_m2: float) -> WeightEstimate:
    m, _ = _rho(code)
    kmin, ktyp, kmax, v_typ, formula = estimate_2d(code, area_m2)
    a = area_m2 * m.fill
    return WeightEstimate(kmin, ktyp, kmax, a * m.t_min, v_typ, a * m.t_max, "2d", formula)


def estimate_count(code: str, n: int = 1) -> WeightEstimate:
    m, spec = _rho(code)
    g = spec.mean_item_g or 13.4
    typ = g / 1000 * n
    lo, hi = COUNT_RANGE
    vol = m.default_vol_m3 * n
    formula = f"크기 정보 없음 → 개수 {n} × 평균무게 {g} g ({spec.item_source or 'Andriolo & Gonçalves 2024'}) (×{lo}–{hi})"
    return WeightEstimate(typ * lo, typ, typ * hi * spec.wet_factor_max, vol * lo, vol, vol * hi, "count", formula)


def apply_estimate(o: LitterObject, e: WeightEstimate) -> LitterObject:
    o.kg_min, o.kg_typ, o.kg_max = e.kg_min, e.kg_typ, e.kg_max
    o.vol_min_m3, o.volume_m3, o.vol_max_m3 = e.vol_min, e.vol_typ, e.vol_max
    o.basis = e.basis; o.note = e.formula
    o.item_kg_max = e.kg_max
    return o


def fill_proxies(objs: list[LitterObject]) -> int:
    """부피 측정이 실패한 물체(basis='proxy' 이고 volume 0) 에 같은 종류의 측정 부피 중앙값을 넣는다. 반환: 대체한 개수."""
    n = 0
    by_code: dict[str, list[tuple[int, float]]] = {}
    rank = {"높음": 0, "중간": 1, "": 2, "낮음": 3}
    for o in objs:
        if o.basis == "3d" and o.volume_m3 > 0:
            level = next((k for k in ("높음", "중간", "낮음") if o.confidence.startswith(k)), "")
            by_code.setdefault(o.code, []).append((rank[level], o.volume_m3))
    for o in objs:
        if o.basis != "proxy" or o.volume_m3 > 0:
            continue
        cands = sorted(by_code.get(o.code, []))
        if cands:
            best = min(r for r, _ in cands)
            vols = [v for r, v in cands if r <= max(best, 1)]        # 신뢰 높음·중간이 있으면 그것만
            vol = float(np.median(vols)); why = f"같은 종류 측정 부피 {len(vols)}개 중앙값"
        else:
            vol = material(o.code).default_vol_m3; why = "종류별 기본 부피(가정값)"
        apply_estimate(o, estimate_proxy(o.code, vol, why)); n += 1
    return n


def _shoelace(pts) -> float:
    s = 0.0
    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
        s += x1 * y2 - x2 * y1
    return abs(s) / 2


def load_geojson(path: str | Path, crs_m: str = "EPSG:5186") -> list[LitterObject]:
    """업체 GeoJSON 라벨 (문갑도 형식: Polygon + material_code, area_sqm, weight_kg, center_lon/lat, detection_seq, image_path)."""
    from pyproj import Transformer
    tr = Transformer.from_crs("EPSG:4326", crs_m, always_xy=True)
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    feats = d["features"] if d.get("type") == "FeatureCollection" else [d]
    out = []
    for i, f in enumerate(feats):
        p = f.get("properties", {})
        g = f.get("geometry", {})
        ring = g["coordinates"][0] if g.get("type") == "Polygon" else None
        if ring:
            if ring[0] == ring[-1]:
                ring = ring[:-1]
            xy = [tr.transform(lo, la) for lo, la in ring]
            xs, ys = [q[0] for q in xy], [q[1] for q in xy]
            w, h = max(xs) - min(xs), max(ys) - min(ys)
            area_poly = _shoelace(xy)
        else:
            w = h = area_poly = 0.0
        lon = p.get("center_lon") if p.get("center_lon") is not None else (sum(q[0] for q in ring) / len(ring))
        lat = p.get("center_lat") if p.get("center_lat") is not None else (sum(q[1] for q in ring) / len(ring))
        x, y = tr.transform(lon, lat)
        raw = str(p.get("material_code", p.get("class", "UNK")))
        m = material(raw)
        area = float(p.get("area_sqm") or area_poly or 0.0)
        seq = int(p.get("detection_seq", i + 1))
        region = p.get("region", "OBJ")
        o = LitterObject(obj_id=f"{region}_{seq:04d}_{raw}", seq=seq, code=m.code, class_name=m.class_name, class_ko=m.ko,
                         lon=float(lon), lat=float(lat), x_m=float(x), y_m=float(y), area_m2=area, w_m=round(w, 3), h_m=round(h, 3),
                         company_kg=(float(p["weight_kg"]) if p.get("weight_kg") is not None else None),
                         kg_min=0, kg_typ=0, kg_max=0, volume_m3=0, image_path=p.get("image_path"),
                         survey_date=p.get("survey_date") or p.get("captured_at"), source=f"업체 라벨 {raw}")
        apply_estimate(o, estimate_2d_full(raw, area))
        out.append(o)
    out.sort(key=lambda o: o.seq)
    return out


def set_weight_source(objs: list[LitterObject], source: str, stat: str = "typ", calib: dict | None = None) -> None:
    """계획 무게 결정. calib = {클래스: 배수} 는 우리 추정값에만 곱한다 (실측 보정)."""
    calib = calib or {}
    for o in objs:
        if source == "company" and o.company_kg is not None:
            o._plan_kg = o.company_kg
            item_kg = o.company_kg
        elif o.kg_by:      # 혼합 정거장: 종류별로 보정
            idx = {"min": 0, "max": 2}.get(stat, 1)
            o._plan_kg = sum(v[idx] * float(calib.get(c, 1.0)) for c, v in o.kg_by.items())
            item_kg = o.item_kg_max * float(max((calib.get(c, 1.0) for c in o.kg_by), default=1.0))
        else:
            f = float(calib.get(o.code, 1.0))
            o._plan_kg = {"min": o.kg_min, "max": o.kg_max}.get(stat, o.kg_typ) * f
            item_kg = o.plan_kg
        # 대표값이 23 kg 를 넘거나, 큰 물체(3 m² 이상) 인데 최대값이 23 kg 를 넘으면 2인 운반 대상
        if o.n_items > 1:
            o.heavy = o.item_kg_max > NIOSH_LIFT_LIMIT_KG
        else:
            o.heavy = item_kg > NIOSH_LIFT_LIMIT_KG or (o.area_m2 >= 3.0 and o.kg_max > NIOSH_LIFT_LIMIT_KG)


# ─────────────────────────── 구역 묶기 ───────────────────────────
def cluster_by_matrix(D: np.ndarray, link_m: float) -> list[int]:
    """거리 행렬로 single-linkage (union-find). 반환: 각 노드의 구역 번호 (등장 순 0..)."""
    n = D.shape[0]
    parent = list(range(n))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]; i = parent[i]
        return i
    for i in range(n):
        for j in range(i + 1, n):
            if D[i, j] <= link_m:
                a, b = find(i), find(j)
                if a != b:
                    parent[a] = b
    remap: dict[int, int] = {}
    return [remap.setdefault(find(i), len(remap)) for i in range(n)]


COMPASS = ["북", "북동", "동", "남동", "남", "남서", "서", "북서"]


def compass_name(dx: float, dy: float) -> str:
    ang = (math.degrees(math.atan2(dx, dy)) + 360) % 360
    return COMPASS[int((ang + 22.5) // 45) % 8]


# ─────────────────────────── 순회 ───────────────────────────
def tour_order(D, start: int, nodes: list[int], round_trip: bool = True) -> list[int]:
    """start 에서 출발해 nodes 를 모두 도는 순서 (거리 행렬 D). 8개 이하 완전탐색, 그 외 NN + 2-opt + or-opt."""
    n = len(nodes)
    if n == 0:
        return []

    def L(order):
        s, pos = 0.0, start
        for i in order:
            s += D[pos, i]; pos = i
        return s + (D[pos, start] if round_trip else 0.0)
    if n <= 8:
        from itertools import permutations
        best, best_o = float("inf"), list(nodes)
        for perm in permutations(nodes):
            if round_trip and n > 1 and perm[0] > perm[-1]:
                continue
            l = L(perm)
            if l < best:
                best, best_o = l, list(perm)
        return best_o
    rem, pos, order = set(nodes), start, []
    while rem:
        nxt = min(rem, key=lambda i: D[pos, i]); order.append(nxt); rem.remove(nxt); pos = nxt
    best = L(order); improved = True
    while improved:
        improved = False
        for i in range(n - 1):
            for j in range(i + 1, n):
                cand = order[:i] + order[i:j + 1][::-1] + order[j + 1:]
                l = L(cand)
                if l < best - 1e-9:
                    order, best, improved = cand, l, True
        for i in range(n):
            for j in range(n):
                if i == j:
                    continue
                cand = order[:]; v = cand.pop(i); cand.insert(j, v)
                l = L(cand)
                if l < best - 1e-9:
                    order, best, improved = cand, l, True
    return order


# ─────────────────────────── 계획 ───────────────────────────
@dataclass
class Zone:
    zone_id: int
    step: int
    name: str
    cx_m: float
    cy_m: float
    lon: float
    lat: float
    objects: list[str]
    n: int
    n_items: int
    by_code: dict[str, int]
    kg_plan: float
    kg_min: float
    kg_typ: float
    kg_max: float
    company_kg: float
    volume_m3: float
    bags: int
    bag_limit: str
    heavy_ids: list[str]
    tools: list[str]
    notes: list[str]
    path_lonlat: list[tuple[float, float]]       # 이 구역을 도는 동안의 전체 경로 (직전 위치 → … → 마지막 물체, 복귀 포함)
    dist_from_prev_m: float                      # 직전 위치 → 구역 첫 물체 (환산 m)
    dist_total_m: float                          # 이 구역에서 걸은 전체 환산 m (복귀 포함)
    walk_min: float
    work_min: float
    cum_min: float
    cum_bags: int
    returns: int = 0                             # 들고 이동 모드에서 출발지 복귀 횟수
    day: int = 1
    team: int = 1


@dataclass
class CollectPlan:
    site: str
    survey_date: str
    depot: dict
    n_objects: int
    n_skipped: int
    totals: dict
    zones: list[Zone]
    route_lonlat: list[tuple[float, float]]
    route_len_m: float
    total_walk_min: float
    total_work_min: float
    total_min: float
    days: list[dict]
    teams: list[dict]
    equipment: list[str]
    by_code: dict[str, dict]
    params: dict
    terrain_used: bool
    terrain_desc: str
    assumptions: list[str]
    objects: list[LitterObject]

    def to_dict(self) -> dict:
        d = {k: v for k, v in asdict(self).items() if k not in ("objects", "zones")}
        d["zones"] = [asdict(z) for z in self.zones]
        d["objects"] = [o.to_dict() for o in self.objects]
        return d


def _bags(kg: float, m3: float, p: PlanParams) -> tuple[int, str]:
    by_w, by_v = kg / p.bag_kg, m3 * p.bulk_factor / p.bag_m3
    if by_v >= by_w:
        return max(math.ceil(by_v), 1 if m3 > 0 or kg > 0 else 0), "volume"
    return max(math.ceil(by_w), 1), "weight"


def make_collect_plan(objs_all: list[LitterObject], params: CollectParams | None = None, *, site: str = "",
                      center_xy: tuple[float, float] | None = None, crs_m: str = "EPSG:5186",
                      terrain=None) -> CollectPlan:
    p = params or CollectParams()
    from pyproj import Transformer
    to_m = Transformer.from_crs("EPSG:4326", crs_m, always_xy=True)
    to_ll = Transformer.from_crs(crs_m, "EPSG:4326", always_xy=True)

    set_weight_source(objs_all, p.weight_source, p.weight_stat, p.calib)
    for o in objs_all:
        o.included = (p.include_codes is None or o.code in p.include_codes) and o.plan_kg >= p.min_kg
        o.zone, o.order = -1, 0
    objs = [o for o in objs_all if o.included]

    # 출발지
    if p.depot_lonlat is not None:
        depot_xy = to_m.transform(*p.depot_lonlat); depot_ll = p.depot_lonlat
    elif objs_all:
        depot_xy = (float(np.mean([o.x_m for o in objs_all])), float(np.mean([o.y_m for o in objs_all])))
        depot_ll = to_ll.transform(*depot_xy)
    else:
        depot_xy, depot_ll = (0.0, 0.0), (0.0, 0.0)
    if center_xy is None:
        center_xy = (float(np.mean([o.x_m for o in objs_all])), float(np.mean([o.y_m for o in objs_all]))) if objs_all else depot_xy

    # 거리 행렬: 0 = 출발지, 1.. = 물체
    nodes_xy = [depot_xy] + [(o.x_m, o.y_m) for o in objs]
    n = len(objs)
    mode = p.travel if p.travel in ("walk", "boat") else "walk"
    if terrain is not None and n:
        terrain.force_passable(nodes_xy)
        D = terrain.matrix(nodes_xy, mode)
        unreachable = ~np.isfinite(D)
        if unreachable.any():      # 도달 불가 쌍은 직선 × 우회 × 3 로 대체 (경고)
            SL = np.array([[math.hypot(a[0] - b[0], a[1] - b[1]) for b in nodes_xy] for a in nodes_xy])
            D = np.where(unreachable, SL * p.detour * 3, D)

        def seg(i, j):
            return terrain.path(nodes_xy[i], nodes_xy[j], mode)
    else:
        D = np.array([[math.hypot(a[0] - b[0], a[1] - b[1]) * p.detour for b in nodes_xy] for a in nodes_xy])

        def seg(i, j):
            return [nodes_xy[i], nodes_xy[j]]

    # 구역
    labels = cluster_by_matrix(D[1:, 1:], p.link_m) if n else []
    for o, l in zip(objs, labels):
        o.zone = l
    groups: dict[int, list[int]] = {}
    for k, o in enumerate(objs):
        groups.setdefault(o.zone, []).append(k + 1)
    zone_ids = sorted(groups)
    zone_kg = {z: sum(objs[i - 1].plan_kg for i in groups[z]) for z in zone_ids}
    zone_work = {z: sum(p.item_min * objs[i - 1].n_items + p.min_per_m2 * objs[i - 1].area_m2 for i in groups[z]) / max(p.workers, 1)
                 + p.heavy_extra_min * sum(1 for i in groups[z] if objs[i - 1].heavy) for z in zone_ids}

    # 구역 사이 거리 = 가장 가까운 물체 쌍, 출발지 → 구역 = 가장 가까운 물체
    Z = len(zone_ids)
    ZD = np.zeros((Z + 1, Z + 1))
    for a, za in enumerate(zone_ids, 1):
        ZD[0, a] = ZD[a, 0] = min(D[0, i] for i in groups[za])
        for b, zb in enumerate(zone_ids, 1):
            if a != b:
                ZD[a, b] = min(D[i, j] for i in groups[za] for j in groups[zb])

    walk_min_per_m = 60.0 / (p.walk_kmh * 1000)
    if p.objective == "weight" and Z:
        rem, pos, order = list(range(1, Z + 1)), 0, []
        while rem:
            nxt = max(rem, key=lambda a: zone_kg[zone_ids[a - 1]] / max(ZD[pos, a] * walk_min_per_m + zone_work[zone_ids[a - 1]], 1e-6))
            order.append(nxt); rem.remove(nxt); pos = nxt
    else:
        order = tour_order(ZD, 0, list(range(1, Z + 1)), p.round_trip)
    ordered = [zone_ids[a - 1] for a in order]

    # 이름
    cents = {z: (float(np.mean([objs[i - 1].x_m for i in groups[z]])), float(np.mean([objs[i - 1].y_m for i in groups[z]]))) for z in zone_ids}
    dir_count: dict[str, int] = {}
    for z in ordered:
        d = compass_name(cents[z][0] - center_xy[0], cents[z][1] - center_xy[1]); dir_count[d] = dir_count.get(d, 0) + 1
    seen: dict[str, int] = {}; names: dict[int, str] = {}
    circ = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
    for z in ordered:
        d = compass_name(cents[z][0] - center_xy[0], cents[z][1] - center_xy[1]); seen[d] = seen.get(d, 0) + 1
        names[z] = f"{d}쪽 {p.zone_word}" + (f" {circ[seen[d] - 1]}" if dir_count[d] > 1 and seen[d] - 1 < len(circ) else (f" {seen[d]}" if dir_count[d] > 1 else ""))

    # 순회 실행 (적재량·복귀 포함) — 한 팀이 seq 순서로 구역을 도는 함수
    cap_kg = p.workers * p.carry_kg_per_person
    cap_bags = p.workers * p.carry_bags_per_person
    route_pts: list[tuple[float, float]] = [depot_xy]

    def traverse(seq: list[int], start_step: int, team: int) -> tuple[list[Zone], float, float]:
        """반환 (zones, 복귀 환산 m, 복귀 분). 출발지에서 시작해 seq 구역을 차례로 돈다."""
        zs: list[Zone] = []
        pos, cum_min, cum_bags = 0, 0.0, 0
        load = {"kg": 0.0, "bags": 0.0}

        def speed_factor():
            return 1.0 - min(0.4, p.load_slow * load["kg"] / max(p.workers, 1))

        for k0, z in enumerate(seq):
            step = start_step + k0
            members = groups[z][:]
            inner = []
            cur = pos
            while members:
                nxt = min(members, key=lambda i: D[cur, i]); inner.append(nxt); members.remove(nxt); cur = nxt
            for k, i in enumerate(inner, 1):
                objs[i - 1].order = k
            d_in = D[pos, inner[0]]
            segs: list[list[tuple[float, float]]] = []
            walk_min = 0.0; dist_total = 0.0; returns = 0
            cur = pos
            for i in inner:
                o = objs[i - 1]
                o_bags = _bags(o.plan_kg, o.volume_m3, p.bag)[0]
                if p.carry == "carry" and (load["kg"] > 0 or load["bags"] > 0) and (load["kg"] + o.plan_kg > cap_kg or load["bags"] + o_bags > cap_bags):
                    back = D[cur, 0]; out = D[0, i]
                    walk_min += back * walk_min_per_m / speed_factor()
                    load["kg"], load["bags"] = 0.0, 0.0
                    walk_min += out * walk_min_per_m / speed_factor()
                    dist_total += back + out
                    segs.append(seg(cur, 0)); segs.append(seg(0, i)); returns += 1
                else:
                    d = D[cur, i]
                    walk_min += d * walk_min_per_m / speed_factor(); dist_total += d
                    segs.append(seg(cur, i))
                if p.carry == "carry":
                    load["kg"] += o.plan_kg; load["bags"] += o_bags
                cur = i
            pos = cur
            L = [objs[i - 1] for i in inner]
            kg_plan = sum(o.plan_kg for o in L); m3 = sum(o.volume_m3 for o in L)
            bags, lim = _bags(kg_plan, m3, p.bag)
            heavy = [o.obj_id for o in L if o.heavy]
            work_min = zone_work[z]
            cum_min += walk_min + work_min; cum_bags += bags
            by_code: dict[str, int] = {}
            for o in L:
                if o.kg_by:
                    for c, v in o.kg_by.items():
                        by_code[c] = by_code.get(c, 0) + int(v[5] if len(v) > 5 else 1)
                else:
                    by_code[o.code] = by_code.get(o.code, 0) + 1
            tools = sorted({material(c).tool for c in by_code if material(c).tool})
            notes = []
            if heavy:
                notes.append(f"무거운 물체 {len(heavy)}개 → 2인 이상 또는 장비 (NIOSH 23 kg 초과 가능)")
            big = [o for o in L if o.area_m2 >= 3.0 and o.n_items == 1]
            if big:
                notes.append(f"면적 3 m² 이상 큰 물체 {len(big)}개 (" + ", ".join(f"{o.class_ko} {o.area_m2:.1f} m²" for o in big) + ")")
            if returns:
                notes.append(f"적재량 초과로 출발지 복귀 {returns}회 포함")
            path_xy = [q for sg in segs for q in sg]
            route_pts.extend(path_xy[1:] if path_xy else [])
            lon, lat = to_ll.transform(*cents[z])
            zs.append(Zone(zone_id=z, step=step, name=names[z], cx_m=cents[z][0], cy_m=cents[z][1], lon=lon, lat=lat,
                           objects=[o.obj_id for o in L], n=len(L), n_items=sum(o.n_items for o in L), by_code=by_code, kg_plan=kg_plan,
                           kg_min=sum(o.kg_min for o in L), kg_typ=sum(o.kg_typ for o in L), kg_max=sum(o.kg_max for o in L),
                           company_kg=sum(o.company_kg or 0 for o in L), volume_m3=m3, bags=bags, bag_limit=lim,
                           heavy_ids=heavy, tools=tools, notes=notes,
                           path_lonlat=[to_ll.transform(x, y) for x, y in path_xy], dist_from_prev_m=float(d_in),
                           dist_total_m=float(dist_total), walk_min=walk_min, work_min=work_min, cum_min=cum_min,
                           cum_bags=cum_bags, returns=returns, team=team))
        back_m, back_min = 0.0, 0.0
        if p.round_trip and zs:
            back_m = float(D[pos, 0]); back_min = back_m * walk_min_per_m / speed_factor()
            route_pts.extend(seg(pos, 0)[1:])
        return zs, back_m, back_min

    # 팀 분할: 한 팀 순회 결과를 시간 균형으로 연속 구간 T개로 나눔 (각 팀은 출발지에서 자기 구간만)
    n_teams = max(int(p.teams), 1)
    if n_teams > 1 and len(ordered) > 1:
        single, _, _ = traverse(ordered, 1, 1)
        route_pts[:] = [depot_xy]
        tot = sum(z.walk_min + z.work_min for z in single)
        segments: list[list[int]] = []; cur_seg: list[int] = []; cum = 0.0
        for k, z in enumerate(single):
            cur_seg.append(ordered[k]); cum += z.walk_min + z.work_min
            remaining = len(single) - k - 1
            teams_left = n_teams - len(segments) - 1          # 아직 구역을 못 받은 뒤 팀 수
            if len(segments) < n_teams - 1 and remaining >= teams_left and (cum >= (len(segments) + 1) * tot / n_teams or remaining == teams_left):
                segments.append(cur_seg); cur_seg = []
        if cur_seg:
            segments.append(cur_seg)
    else:
        segments = [ordered] if ordered else []

    zones: list[Zone] = []
    team_info: list[dict] = []
    days: list[dict] = []
    route_len = 0.0; total_walk = 0.0; total_work = 0.0
    day_cap = p.hours_per_day * 60
    step0 = 1
    for t, seq in enumerate(segments, 1):
        zs, back_m, back_min = traverse(seq, step0, t)
        step0 += len(zs)
        walk = sum(z.walk_min for z in zs) + back_min; work = sum(z.work_min for z in zs)
        route_len += sum(z.dist_total_m for z in zs) + back_m
        total_walk += walk; total_work += work
        day, acc = 1, 0.0
        for z in zs:
            if acc > 0 and acc + z.walk_min + z.work_min > day_cap:
                days.append({"team": t, "day": day, "steps": [zz.step for zz in zs if zz.day == day and zz.team == t], "minutes": round(acc, 1)})
                day, acc = day + 1, 0.0
            z.day = day; acc += z.walk_min + z.work_min
        if zs:
            days.append({"team": t, "day": day, "steps": [zz.step for zz in zs if zz.day == day], "minutes": round(acc + back_min, 1)})
        team_info.append({"team": t, "zones": [z.step for z in zs], "n": sum(z.n for z in zs), "kg": sum(z.kg_plan for z in zs),
                          "bags": sum(z.bags for z in zs), "walk_min": walk, "work_min": work, "minutes": walk + work,
                          "days": day if zs else 0, "route_m": sum(z.dist_total_m for z in zs) + back_m})
        zones.extend(zs)
    total_min = max((ti["minutes"] for ti in team_info), default=0.0)     # 팀이 동시에 일하므로 가장 오래 걸리는 팀 기준

    kg_plan = sum(o.plan_kg for o in objs)
    vol = sum(o.volume_m3 for o in objs)
    basis_counts: dict[str, int] = {}
    for o in objs_all:
        basis_counts[o.basis] = basis_counts.get(o.basis, 0) + 1
    truth = [o for o in objs_all if o.truth_m3]
    totals = {
        "kg_plan": kg_plan, "kg_min": sum(o.kg_min for o in objs), "kg_typ": sum(o.kg_typ for o in objs),
        "kg_max": sum(o.kg_max for o in objs), "company_kg": sum(o.company_kg or 0 for o in objs),
        "volume_m3": vol, "area_m2": sum(o.area_m2 for o in objs), "items": sum(o.n_items for o in objs),
        "bags": sum(z.bags for z in zones), "heavy": sum(len(z.heavy_ids) for z in zones), "zones": len(zones),
        "returns": sum(z.returns for z in zones), "teams": len(team_info), "work_min_sum": total_walk + total_work,
        "tonbags": math.ceil(max(kg_plan / p.bag.tonbag_kg, vol * p.bag.bulk_factor / p.bag.tonbag_m3)) if objs else 0,
        "basis_counts": basis_counts,
        "truth": ({"n": len(truth), "measured_m3": sum(o.volume_m3 for o in truth if o.basis == "3d"),
                   "truth_m3": sum(o.truth_m3 for o in truth if o.basis == "3d"),
                   "n_measured": sum(1 for o in truth if o.basis == "3d")} if truth else None),
    }
    by_code: dict[str, dict] = {}
    for o in objs:
        parts = o.kg_by.items() if o.kg_by else [(o.code, [o.kg_min, o.kg_typ, o.kg_max, o.volume_m3, o.area_m2, 1])]
        for c, v in parts:
            m = material(c)
            d = by_code.setdefault(m.code, {"ko": m.ko, "count": 0, "area_m2": 0.0, "kg_plan": 0.0, "kg_min": 0.0,
                                            "kg_typ": 0.0, "kg_max": 0.0, "company_kg": 0.0, "volume_m3": 0.0,
                                            "color": m.color, "handling": m.handling_note, "tool": m.tool, "basis": {}})
            cnt = int(v[5]) if len(v) > 5 else 1
            d["count"] += cnt; d["area_m2"] += v[4] if len(v) > 4 else o.area_m2
            d["kg_min"] += v[0]; d["kg_typ"] += v[1]; d["kg_max"] += v[2]; d["volume_m3"] += v[3] if len(v) > 3 else o.volume_m3
            d["basis"][o.basis] = d["basis"].get(o.basis, 0) + cnt
        d_main = by_code[material(o.code).code]
        d_main["kg_plan"] += o.plan_kg; d_main["company_kg"] += o.company_kg or 0

    equipment = [f"마대 {math.ceil(totals['bags'] * 1.2)}장 (계산 {totals['bags']}장 + 여유 20 %)", "장갑·집게 인원수만큼"]
    tools = sorted({t for z in zones for t in z.tools})
    if tools:
        equipment.append(" / ".join(tools) + " (로프·그물 자르기 등)")
    if totals["tonbags"] >= 1 and vol * p.bag.bulk_factor > 0.5:
        equipment.append(f"톤백 {totals['tonbags']}개 또는 집결지 적재 공간 {vol * p.bag.bulk_factor:.1f} m³")
    if totals["heavy"]:
        equipment.append(f"무거운 물체 {totals['heavy']}개 → 2인 운반 또는 손수레")
    if p.travel == "boat":
        equipment.append("보트·구명조끼, 승·하선 지점 사전 확인")
    equipment.append("식수·구급약" + (", 물때표 확인 (갯바위 구간)" if p.zone_word == "해안" else ""))

    survey = sorted({o.survey_date for o in objs_all if o.survey_date})
    assumptions = []
    if basis_counts.get("3d"):
        assumptions.append(f"무게 = 드론 영상에서 3D 복원한 겉 부피(높이지도) × 종류별 겉보기 밀도 (classes.py; 스티로폼은 출처 있음, 나머지는 가정값). "
                           f"부피 측정 신뢰도에 따라 범위를 ×0.8–1.25(높음)·×0.6–1.5(중간)·×0.5–2(낮음) 로 넓힘 (가정값)")
    if basis_counts.get("proxy"):
        assumptions.append(f"부피 측정이 실패한 물체 {basis_counts['proxy']}개는 같은 종류의 측정 부피 중앙값(없으면 종류별 기본 부피)으로 대체 (범위 ×0.5–2, 가정값)")
    if basis_counts.get("2d") or basis_counts.get("agg"):
        assumptions.append("면적만 있는 물체(라벨·항공 정사영상)는 면적 × 채움률 × 두께 × 겉보기 밀도 로 추정 (채움률·두께는 가정값)")
    if basis_counts.get("count"):
        assumptions.append("크기 정보가 없는 탐지는 개수 × 종류별 평균무게 (Andriolo 2024) 로 추정")
    if any(o.company_kg is not None for o in objs_all):
        assumptions.append("'업체 방식' 무게 = 면적 × 재질별 고정계수(스티로폼 0.012, 로프·어구 0.024, 플라스틱 0.020 kg/m²) — 업체 라벨의 weight_kg 와 같은 방식, 실측 아님 (비교용)")
    assumptions += [
        (f"이동 거리 = {terrain.describe()} 위 최단경로의 걷기 환산 m (숲·절벽 ×3, 물 통행 불가, 보트 모드 물 ×0.5, 승·하선 100 m 상당; 전부 가정값)"
         if terrain is not None else "이동 거리 = 직선거리 × 우회 배수(가정값). 건물·갯바위·절벽 구간은 실제로 더 걸릴 수 있음"),
        "작업 시간 = 물체당 기본 시간 + 면적 비례 시간 (가정값) ÷ 인원, 짐이 늘면 걷기 감속. 날씨·물때 미반영",
        "들고 이동 모드: 팀 적재량(인원 × 1인 kg·마대 수) 을 넘으면 출발지에 내려놓고 돌아옴 (가정값)",
        "1인 들기 한계 23 kg 는 NIOSH Revised Lifting Equation",
        "검출 누락은 반영하지 않은 최소 추정치 (현장에서 탐지에 없는 쓰레기도 함께 수거)",
    ] + (["여러 팀: 한 팀 순회를 시간 균형으로 연속 구간으로 나눠 각 팀이 출발지에서 자기 구간만 돈다. 총 시간은 가장 오래 걸리는 팀 기준"] if n_teams > 1 else []) \
      + ([f"실측 보정 계수 적용: " + ", ".join(f"{material(c).ko} ×{v:.2f}" for c, v in (p.calib or {}).items())] if p.calib else [])
    return CollectPlan(
        site=site, survey_date=", ".join(survey),
        depot={"lon": depot_ll[0], "lat": depot_ll[1], "x_m": depot_xy[0], "y_m": depot_xy[1], "name": p.depot_name},
        n_objects=len(objs), n_skipped=len(objs_all) - len(objs), totals=totals, zones=zones,
        route_lonlat=[to_ll.transform(x, y) for x, y in route_pts], route_len_m=route_len, total_walk_min=total_walk,
        total_work_min=total_work, total_min=total_min, days=days, teams=team_info, equipment=equipment, by_code=by_code,
        params=p.to_dict(), terrain_used=terrain is not None, terrain_desc=terrain.describe() if terrain is not None else "",
        assumptions=assumptions, objects=objs_all)


# ─────────────────────────── 저장 ───────────────────────────
ZONE_COLS = ["step", "team", "day", "name", "n", "n_items", "kg_plan", "kg_min", "kg_max", "company_kg", "volume_m3", "bags",
             "dist_from_prev_m", "dist_total_m", "walk_min", "work_min", "cum_min", "cum_bags", "returns", "lon", "lat", "tools", "notes"]
ZONE_KO = ["순서", "팀", "일차", "구역", "개수", "물체 수", "계획 무게(kg)", "최소(kg)", "최대(kg)", "업체 방식(kg)", "부피(m³)", "마대(장)",
           "접근(m)", "구역 이동 합(m)", "이동(분)", "작업(분)", "누적(분)", "누적 마대", "복귀 횟수", "경도", "위도", "도구", "주의"]
OBJ_COLS = ["step", "order", "obj_id", "class_ko", "included", "basis_ko", "dims", "area_m2", "volume_m3", "vol_min_m3", "vol_max_m3",
            "truth_m3", "confidence", "plan_kg", "kg_min", "kg_typ", "kg_max", "company_kg", "heavy", "n_items", "lon", "lat", "source", "image_path"]
OBJ_KO = ["구역 순서", "구역 내 순서", "ID", "종류", "수거 대상", "무게 근거", "크기", "면적(m²)", "부피(m³)", "부피 최소(m³)", "부피 최대(m³)",
          "정답 부피(m³)", "측정 신뢰도", "계획 무게(kg)", "최소(kg)", "대표(kg)", "최대(kg)", "업체 방식(kg)", "2인 운반", "물체 수", "경도", "위도", "출처", "사진"]


def save_plan_files(plan: CollectPlan, out_dir: str | Path) -> dict[str, Path]:
    import csv
    out = Path(out_dir); out.mkdir(parents=True, exist_ok=True)
    files = {}
    (out / "plan.json").write_text(json.dumps(plan.to_dict(), ensure_ascii=False, indent=2, default=float), encoding="utf-8")
    files["plan.json"] = out / "plan.json"
    step_of = {oid: z.step for z in plan.zones for oid in z.objects}

    def zrow(z: Zone):
        d = asdict(z); d["tools"] = ", ".join(z.tools); d["notes"] = " / ".join(z.notes)
        return [round(d[c], 3) if isinstance(d[c], float) else d[c] for c in ZONE_COLS]

    def orow(o: LitterObject):
        d = o.to_dict(); d["step"] = step_of.get(o.obj_id, 0)
        return [round(d[c], 4) if isinstance(d[c], float) else d[c] for c in OBJ_COLS]

    with open(out / "zones.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(ZONE_KO); [w.writerow(zrow(z)) for z in plan.zones]
    objs_sorted = sorted(plan.objects, key=lambda o: (0 if o.included else 1, step_of.get(o.obj_id, 0), o.order))
    with open(out / "objects.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(OBJ_KO); [w.writerow(orow(o)) for o in objs_sorted]
    files["zones.csv"] = out / "zones.csv"; files["objects.csv"] = out / "objects.csv"

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = "요약"
        t = plan.totals
        bc = ", ".join(f"{BASIS_KO.get(k, k)} {v}" for k, v in t.get("basis_counts", {}).items())
        rows = [["항목", "값"], ["현장", plan.site], ["조사일", plan.survey_date], ["수거 대상 개수", plan.n_objects],
                ["제외(종류·무게 필터)", plan.n_skipped], ["무게 근거", bc], ["구역 수", t["zones"]], ["계획 무게(kg)", round(t["kg_plan"], 1)],
                ["무게 범위(kg)", f"{t['kg_min']:.1f} – {t['kg_max']:.1f}"], ["업체 방식 무게 합(kg)", round(t["company_kg"], 3)],
                ["부피(m³)", round(t["volume_m3"], 3)], ["마대(장)", t["bags"]], ["이동 거리(걷기 환산 m)", round(plan.route_len_m)],
                ["총 시간(분, 가장 오래 걸리는 팀)", round(plan.total_min)], ["팀 수", len(plan.teams)], ["일수(최대)", max((t["days"] for t in plan.teams), default=0)], ["인원(팀당)", plan.params["workers"]],
                ["이동 방식", plan.params["travel"]], ["운반 방식", plan.params["carry"]], ["최적화 목표", plan.params["objective"]],
                ["지형 반영", plan.terrain_desc if plan.terrain_used else "아니오(직선×우회)"]]
        for r in rows:
            ws.append(r)
        ws2 = wb.create_sheet("작업순서"); ws2.append(ZONE_KO); [ws2.append(zrow(z)) for z in plan.zones]
        ws3 = wb.create_sheet("물체목록"); ws3.append(OBJ_KO); [ws3.append(orow(o)) for o in objs_sorted]
        ws4 = wb.create_sheet("가정"); ws4.append(["가정과 한계"]); [ws4.append([a]) for a in plan.assumptions]
        for s in (ws, ws2, ws3, ws4):
            for c in s[1]:
                c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7"); c.alignment = Alignment(horizontal="center")
            s.freeze_panes = "A2"
            for col in s.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                s.column_dimensions[col[0].column_letter].width = min(max(10, width * 1.6), 60)
        wb.save(out / "수거계획.xlsx"); files["수거계획.xlsx"] = out / "수거계획.xlsx"
    except Exception as e:  # noqa: BLE001
        print(f"  ⚠ xlsx 저장 실패 ({e}) — csv 만 저장")
    return files


def summary_markdown(plan: CollectPlan) -> str:
    t = plan.totals
    bc = ", ".join(f"{BASIS_KO.get(k, k)} {v}개" for k, v in t.get("basis_counts", {}).items())
    L = [f"# {plan.site or '현장'} 쓰레기 수거 계획", "",
         f"- 조사일 {plan.survey_date} · 수거 대상 {plan.n_objects}개 (제외 {plan.n_skipped}) · 구역 {t['zones']}곳 · 출발지 {plan.depot['name']}",
         f"- 무게 근거: {bc}",
         f"- 계획 무게 **{t['kg_plan']:.1f} kg** (추정 범위 {t['kg_min']:.1f} – {t['kg_max']:.1f} kg, 업체 방식 합 {t['company_kg']:.3f} kg)",
         f"- 부피 {t['volume_m3']:.3f} m³ → 마대 {t['bags']}장, 무거운 물체 {t['heavy']}개",
         f"- 이동 {plan.route_len_m / 1000:.1f} km (걷기 환산, {plan.terrain_desc if plan.terrain_used else '직선×우회'}, {plan.params['travel']}/{plan.params['carry']}) · "
         f"이동 {plan.total_walk_min:.0f}분 + 작업 {plan.total_work_min:.0f}분 (전 팀 합) · 가장 오래 걸리는 팀 {plan.total_min / 60:.1f}시간 "
         f"({len(plan.teams)}팀 × {plan.params['workers']}명) → 최대 {max((t['days'] for t in plan.teams), default=0)}일"]
    if t.get("truth") and t["truth"]["n_measured"]:
        tr = t["truth"]
        L.append(f"- 크기를 아는 시험 물체 {tr['n']}개 중 3D 측정 {tr['n_measured']}개: 측정 부피 합 {tr['measured_m3'] * 1000:.1f} L vs 정답 {tr['truth_m3'] * 1000:.1f} L "
                 f"({(tr['measured_m3'] / tr['truth_m3'] - 1) * 100:+.0f} %)")
    L += ["", "## 작업 순서", "", "| 순서 | 팀 | 일차 | 구역 | 개수 | 구성 | 무게(kg) | 마대 | 접근 | 이동 합 | 작업 | 주의 |", "|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for z in plan.zones:
        comp = ", ".join(f"{material(c).ko} {n}" for c, n in sorted(z.by_code.items(), key=lambda kv: -kv[1]))
        L.append(f"| {z.step} | {z.team} | {z.day} | {z.name} | {z.n} | {comp} | {z.kg_plan:.1f} ({z.kg_min:.1f}–{z.kg_max:.1f}) | {z.bags} | "
                 f"{z.dist_from_prev_m:.0f} m | {z.dist_total_m:.0f} m / {z.walk_min:.0f}분 | {z.work_min:.0f}분 | {' / '.join(z.notes) or '-'} |")
    L += ["", "## 준비물", ""] + [f"- {e}" for e in plan.equipment]
    L += ["", "## 종류별", "", "| 종류 | 개수 | 면적(m²) | 부피(L) | 계획 무게(kg) | 범위(kg) | 업체 방식(kg) | 근거 |", "|---|---|---|---|---|---|---|---|"]
    for c, d in sorted(plan.by_code.items(), key=lambda kv: -kv[1]["kg_plan"]):
        bs = ", ".join(f"{BASIS_KO.get(k, k)} {v}" for k, v in d["basis"].items())
        L.append(f"| {d['ko']} ({c}) | {d['count']} | {d['area_m2']:.1f} | {d['volume_m3'] * 1000:.0f} | {d['kg_plan']:.1f} | {d['kg_min']:.1f}–{d['kg_max']:.1f} | {d['company_kg']:.3f} | {bs} |")
    L += ["", "## 가정과 한계", ""] + [f"- {a}" for a in plan.assumptions]
    return "\n".join(L)
